"""Generate tiny reader test fixtures, not a publication deck or course content."""
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
os.environ.update(TEMP=str(ROOT / '.cache/tmp'), TMP=str(ROOT / '.cache/tmp'))
sys.path.insert(0, str(ROOT / '.venv/Lib/site-packages'))
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt, Inches

samples = ROOT / 'samples'
samples.mkdir(exist_ok=True)
pdf = canvas.Canvas(str(samples / 'optimization.pdf'), pagesize=(595, 842))
pdf.setTitle('Optimization - Reader Verification')
for number, heading, lines in [
    (1, 'Gradient descent', ['Gradient descent', '',
        'Gradient descent moves parameters opposite to the gradient.',
        'Learning rate controls the size of each update.', '',
        'w_new = w - learning_rate * gradient', '',
        'Reading exercise:', 'Select the term above to open the translation toolbar.',
        'Ask a small question, then save a note.']),
    (2, 'Generalization', ['Overfitting', '',
        'Overfitting can reduce performance on unseen data.',
        'Compare training performance with validation performance.', '',
        'Reader verification:', 'Navigate back to page 1 using a saved citation.',
        'Reload the application and check that your notes remain.'])]:
    pdf.setFillColor(HexColor('#296b55'))
    pdf.setFont('Helvetica', 10)
    pdf.drawString(50, 795, 'STUDY LOCAL / READER FIXTURE')
    pdf.setFont('Helvetica-Bold', 28)
    pdf.drawString(50, 735, heading)
    pdf.setFillColor(HexColor('#203831'))
    y = 665
    for index, line in enumerate(lines):
        pdf.setFont('Helvetica-Bold' if index == 0 else 'Helvetica', 18 if index == 0 else 12)
        pdf.drawString(50, y, line)
        y -= 32
    pdf.setFont('Helvetica', 10)
    pdf.setFillColor(HexColor('#738179'))
    pdf.drawString(50, 45, 'Local interaction sample. No AI-generated answer.')
    pdf.drawRightString(545, 45, str(number))
    pdf.showPage()
pdf.save()

image = Image.new('RGB', (900, 1200), 'white')
draw = ImageDraw.Draw(image)
draw.text((80, 140), 'Scanned sample - raster image only', fill='black')
draw.text((80, 210), 'This page must not produce extracted text.', fill='black')
scan = canvas.Canvas(str(samples / 'scan-no-ocr.pdf'), pagesize=(595, 842))
scan.drawImage(ImageReader(image), 0, 0, 595, 842)
scan.save()

presentation = Presentation()
for heading, body, note in [
    ('Gradient descent', 'Gradient descent\nLearning rate\n参数沿梯度反方向更新。', '讲者备注：请区分更新方向和更新步长。'),
    ('Overfitting', 'Overfitting\nTraining data / Validation data\n验证集帮助观察泛化表现。', '已有讲者备注也应被提取。')]:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = heading
    slide.placeholders[1].text = body
    slide.notes_slide.notes_text_frame.text = note
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(28 if shape == slide.shapes.title else 20)
table = presentation.slides[1].shapes.add_table(2, 2, Inches(1), Inches(5.3), Inches(7), Inches(1)).table
table.cell(0, 0).text = 'Training data'
table.cell(0, 1).text = 'Validation data'
table.cell(1, 0).text = 'Fit parameters'
table.cell(1, 1).text = 'Check generalization'
presentation.save(samples / 'lecture.pptx')
print('Generated PDF, PPTX and raster-only fixtures in D:/note/samples.')
