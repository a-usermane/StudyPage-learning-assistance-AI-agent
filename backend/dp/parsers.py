"""Local format adapters. No HTTP or storage responsibilities."""
import zipfile
from pypdf import PdfReader
from pptx import Presentation
from backend.domain.models import ParsedDocument

def decode_text(content):
    if content.startswith((b'\xff\xfe', b'\xfe\xff')):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            if "\x00" in text:
                raise ValueError("文件包含二进制内容，无法作为文本读取。")
            return text
        except UnicodeDecodeError:
            continue
    raise ValueError("文本编码无法识别，请转换为 UTF-8 后重试。")

def extract(path, extension):
    if extension == ".pdf":
        with path.open("rb") as source:
            reader = PdfReader(source)
            if reader.is_encrypted:
                raise ValueError("首版不支持加密 PDF，请先导出未加密副本。")
            if len(reader.pages) > 1000:
                raise ValueError("首版最多支持 1000 页，请拆分 PDF。")
            result = [page.extract_text() or "" for page in reader.pages]
    elif extension == ".pptx":
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 10000 or sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("PPTX 解压后内容过大，请简化资料后重试。")
        slides = Presentation(str(path)).slides
        if len(slides) > 1000:
            raise ValueError("首版最多支持 1000 张幻灯片。")
        result = []
        for slide in slides:
            lines = []
            def shapes_text(shapes):
                for shape in shapes:
                    if shape.has_text_frame:
                        lines.append(shape.text_frame.text)
                    if shape.has_table:
                        lines.extend("  |  ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                    if shape.shape_type == 6:
                        shapes_text(shape.shapes)
            shapes_text(slide.shapes)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    lines.extend(["\n讲者备注", frame.text])
            result.append("\n\n".join(lines))
    else:
        text = decode_text(path.read_bytes()).replace("\r\n", "\n").replace("\r", "\n")
        if len(text) > 2_000_000:
            raise ValueError("文本过长，首版最多读取 200 万字符，请拆分文件。")
        result = [text]
    if not result:
        raise ValueError("文件没有页面或幻灯片。")
    if sum(map(len, result)) > 5_000_000:
        raise ValueError("提取文字过多，请拆分资料。")
    return result

class LocalDocumentParser:
    def parse(self, path, extension):
        pages = extract(path, extension)
        blank = sum(not text.strip() for text in pages)
        warning = f"{blank} 页未提取到文字。图片或扫描页可阅读，但首版不支持 OCR。" if blank and extension == ".pdf" else ""
        if extension == ".pptx":
            warning = "PPTX 仅展示提取的文字、表格和已有讲者备注，不保留原版面。"
        return ParsedDocument(pages, warning)
